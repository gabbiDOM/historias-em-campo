"""Completa os dados que a planilha não traz:

- dados/clubes.json: escudo e estádio de cada clube e logo de cada campeonato (Wikidata/Wikipédia);
- dados/fotos.json: foto de cada jogador do Elenco e das Transferencias (site de notas da EA).

Só busca nomes que ainda não estão nos arquivos. Para buscar de novo, apague a entrada do JSON.
"""
import json
import pathlib
import re
import unicodedata
import urllib.parse
import urllib.request

from openpyxl import load_workbook

UA = {"User-Agent": "HistoriasEmCampo/1.0 (https://gabbidom.github.io/historias-em-campo/)"}
WD = "https://www.wikidata.org/w/api.php?format=json&"
CLUBES = pathlib.Path("dados/clubes.json")
FOTOS = pathlib.Path("dados/fotos.json")
# Valores da coluna Clube das Transferencias que não são clubes.
SEM_CLUBE = {"", "livre", "semclube", "base", "aposentado", "aposentou", "dispensado", "none"}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def get(url, headers=UA):
    return urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=30).read()


def get_json(url):
    return json.loads(get(url))


def load(path, default):
    return json.loads(path.read_text("utf-8")) if path.exists() else default


def save(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", "utf-8")


def ler_planilha(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    abas = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [norm(h) for h in rows[0]]
        abas[norm(ws.title)] = [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    return abas


# ---------- Wikidata ----------
def atual(claims):
    """Valor vigente de uma propriedade: rank preferido, senão sem data de fim (P582), senão o último."""
    if not claims:
        return None
    pref = [c for c in claims if c.get("rank") == "preferred"]
    vigentes = [c for c in claims if "P582" not in c.get("qualifiers", {}) and c.get("rank") != "deprecated"]
    c = (pref or vigentes or claims)[-1]
    return c["mainsnak"].get("datavalue", {}).get("value")


def entidade(nome, filtro):
    q = urllib.parse.quote(f"{nome} {filtro}")
    res = get_json(WD + f"action=query&list=search&srlimit=1&srsearch={q}")["query"]["search"]
    if not res:
        return None, None
    qid = res[0]["title"]
    return qid, get_json(WD + f"action=wbgetentities&props=claims|sitelinks&ids={qid}")["entities"][qid]


def commons(arquivo):
    return "https://commons.wikimedia.org/wiki/Special:FilePath/" + urllib.parse.quote(arquivo.replace(" ", "_")) + "?width=160"


def logo(ent):
    arquivo = atual(ent["claims"].get("P154"))
    if arquivo:
        return commons(arquivo)
    # Sem logo livre no Wikidata: imagem principal do artigo na Wikipédia.
    for site, dom in (("dewiki", "de"), ("enwiki", "en"), ("ptwiki", "pt")):
        titulo = ent.get("sitelinks", {}).get(site, {}).get("title")
        if not titulo:
            continue
        r = get_json(f"https://{dom}.wikipedia.org/w/api.php?format=json&action=query&prop=pageimages"
                     f"&piprop=thumbnail&pithumbsize=160&pilicense=any&titles={urllib.parse.quote(titulo)}")
        for p in r["query"]["pages"].values():
            if p.get("thumbnail"):
                return p["thumbnail"]["source"]
    return None


def rotulo(qid):
    labels = get_json(WD + f"action=wbgetentities&props=labels&languages=de|en|pt&ids={qid}")["entities"][qid]["labels"]
    for lang in ("de", "en", "pt"):
        if lang in labels:
            return labels[lang]["value"]
    return None


def resolver_clube(nome):
    qid, ent = entidade(nome, "haswbstatement:P31=Q476028")
    if not ent:
        return {"wikidata": None, "escudo": None, "estadio": None}
    estadio = atual(ent["claims"].get("P115"))
    return {"wikidata": qid, "escudo": logo(ent), "estadio": rotulo(estadio["id"]) if estadio else None}


def resolver_competicao(nome):
    qid, ent = entidade(nome, "haswbstatement:P641=Q2736")
    return {"wikidata": qid, "logo": logo(ent) if ent else None}


# ---------- EA ----------
def itens_ea(dados):
    achados = []

    def walk(o, d=0):
        if achados or d > 14 or not isinstance(o, (dict, list)):
            return
        if isinstance(o, list) and o and isinstance(o[0], dict) and "overallRating" in o[0]:
            achados.append(o)
            return
        for v in (o.values() if isinstance(o, dict) else o):
            walk(v, d + 1)

    walk(dados)
    return achados[0] if achados else []


def foto_ea(nome, clube_pref):
    html = get("https://www.ea.com/pt-br/games/ea-sports-fc/ratings?search=" + urllib.parse.quote(nome),
               {"User-Agent": "Mozilla/5.0"}).decode("utf-8", "ignore")
    m = re.search(r'<script id="__NEXT_DATA__"[^>]*>(.*?)</script>', html, re.S)
    if not m:
        return ""
    alvo = norm(nome)
    cands = [p for p in itens_ea(json.loads(m.group(1)))
             if alvo in (norm(p.get("commonName")), norm(p.get("lastName")), norm(f'{p.get("firstName")} {p.get("lastName")}'))]
    pref = [p for p in cands if clube_pref and norm(clube_pref) in norm((p.get("team") or {}).get("label"))]
    escolhido = (pref or cands or [None])[0]
    url = (escolhido or {}).get("avatarUrl") or ""
    if url:
        try:
            get(url, {"User-Agent": "Mozilla/5.0"})  # alguns jogadores não têm foto (404)
        except Exception:
            url = ""
    return url


def main():
    # Escudos, estádios e fotos são compartilhados entre todas as séries.
    nomes_clubes, nomes_comps, pedidos = set(), set(), []
    for arquivo in sorted(pathlib.Path("dados/series").glob("*/dados.xlsx")):
        abas = ler_planilha(arquivo)
        cfg = {norm(r.get("chave")): r.get("valor") for r in abas.get("clube", [])}
        meu_clube = str(cfg.get("clube") or "")
        nomes_clubes |= {meu_clube}
        nomes_clubes |= {r.get("adversario") for r in abas.get("jogos", [])}
        nomes_clubes |= {r.get("time") for r in abas.get("classificacao", [])}
        nomes_clubes |= {r.get("clube") for r in abas.get("transferencias", [])}
        nomes_comps |= {str(r.get("competicao")).strip() for r in abas.get("jogos", []) if r.get("competicao")}
        pedidos += [(r.get("jogador"), meu_clube) for r in abas.get("elenco", [])]
        pedidos += [(r.get("jogador"), r.get("clube")) for r in abas.get("transferencias", []) if norm(r.get("tipo")).startswith("cheg")]
    nomes_clubes = {str(n).strip() for n in nomes_clubes if n and norm(n) not in SEM_CLUBE}

    dados = load(CLUBES, {"clubes": {}, "competicoes": {}})
    for nome in sorted(nomes_clubes - set(dados["clubes"])):
        try:
            dados["clubes"][nome] = resolver_clube(nome)
            print("clube:", nome, dados["clubes"][nome])
        except Exception as e:  # rede instável: tenta de novo na próxima execução
            print("falhou clube:", nome, e)
    for nome in sorted(nomes_comps - set(dados["competicoes"])):
        try:
            dados["competicoes"][nome] = resolver_competicao(nome)
            print("campeonato:", nome, dados["competicoes"][nome])
        except Exception as e:
            print("falhou campeonato:", nome, e)
    save(CLUBES, dados)

    fotos = load(FOTOS, {})
    for nome, clube in pedidos:
        nome = str(nome or "").strip()
        if not nome or nome in fotos:
            continue
        try:
            fotos[nome] = foto_ea(nome, clube)
            print("foto:", nome, fotos[nome] or "(sem foto)")
        except Exception as e:
            print("falhou foto:", nome, e)
    save(FOTOS, fotos)


if __name__ == "__main__":
    main()
